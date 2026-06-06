import asyncio
from pathlib import Path
from typing import Dict, List

import openpyxl
from playwright.async_api import (
    async_playwright,
    Error as PlaywrightError,
    TimeoutError as PlaywrightTimeoutError,
)

APP_URL = "https://test-auto-sandy.vercel.app/"
MAX_RETRIES = 5

EXPECTED_COLUMNS = ["Mã nhân viên", "Họ tên", "Legal Entity", "Ngày"]
LEGAL_ENTITY_OPTIONS = ["HO", "Branch 1", "Branch 2", "Branch 3"]
ENTITY_TYPE = "Customer"


def read_excel_rows(file_path: Path) -> List[Dict[str, str]]:
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    sheet = workbook.active
    headers = [str(cell).strip() if cell is not None else "" for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if all(cell is None for cell in row):
            continue
        rows.append({header: str(value).strip() if value is not None else "" for header, value in zip(headers, row)})
    return rows


def read_excel_preview(file_path: Path, max_rows: int = 10) -> List[Dict[str, str]]:
    rows = read_excel_rows(file_path)
    return rows[:max_rows]


async def wait_for_text(page, text: str, timeout: int = 5000):
    await page.wait_for_selector(f"text={text}", timeout=timeout)


async def fill_field(page, label_text: str, value: str):
    locator = page.locator(f"xpath=//label[contains(normalize-space(string(.)), '{label_text}')]/following-sibling::input[1]")
    if await locator.count() == 0:
        locator = page.locator(f"xpath=//label[contains(normalize-space(string(.)), '{label_text}')]/../input[1]")
    await locator.fill(value)


async def select_field(page, label_text: str, option_text: str, selector: str = None):
    if selector:
        locator = page.locator(selector)
    else:
        locator = page.locator(f"xpath=//label[contains(normalize-space(string(.)), '{label_text}')]/following-sibling::select[1]")
        if await locator.count() == 0:
            locator = page.locator(f"xpath=//label[contains(normalize-space(string(.)), '{label_text}')]/../select[1]")
    await locator.select_option(label=option_text)


async def select_custom_dropdown(page, dropdown_selector: str, option_text: str):
    await page.click(dropdown_selector)
    await page.wait_for_selector(f"text={option_text}", timeout=5000)
    await page.click(f"text={option_text}")


async def process_row(page, idx: int, row_data: Dict[str, str], logs: List[str], entity_type: str) -> Dict[str, str]:
    row_status = {"row": idx, "success": False, "error": ""}
    code_value = row_data.get("Mã nhân viên", "")

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            log_message = f"Dòng {idx}, thử lần {attempt} (Loại {entity_type})"
            print(log_message)
            logs.append(log_message)

            await page.goto(APP_URL, wait_until="networkidle")
            await page.select_option("select", label=entity_type)

            await fill_field(page, "Code", code_value)
            await select_field(page, "Legal Entity", row_data.get("Legal Entity", LEGAL_ENTITY_OPTIONS[0]))
            await fill_field(page, "Text / Name", row_data.get("Họ tên", ""))
            await fill_field(page, "Ngày", row_data.get("Ngày", ""))

            await page.click("text=Tạo")
            await page.wait_for_timeout(1000)
            await wait_for_text(page, code_value, timeout=5000)

            success_message = f"Dòng {idx} tạo thành công"
            print(success_message)
            logs.append(success_message)
            row_status["success"] = True
            return row_status

        except (PlaywrightTimeoutError, PlaywrightError) as exc:
            error_message = f"Lần thử {attempt} thất bại: {exc}"
            print(error_message)
            logs.append(error_message)
            row_status["error"] = str(exc)
            if attempt < MAX_RETRIES:
                await page.reload()
                await page.wait_for_load_state("networkidle")
            else:
                screenshot_file = Path(f"loi_dong_{idx}.png")
                await page.screenshot(path=str(screenshot_file), full_page=True)
                failure_message = f"Dòng {idx} lỗi sau {MAX_RETRIES} lần. Screenshot: {screenshot_file}"
                print(failure_message)
                logs.append(failure_message)
                row_status["error"] = f"Failed after {MAX_RETRIES} attempts"
                return row_status


async def run_import(file_path: Path, entity_type: str):
    rows = read_excel_rows(file_path)
    if not rows:
        return {"message": "Không có dữ liệu trong file Excel.", "failed_rows": [], "logs": []}

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=False)
        page = await browser.new_page()
        results = []
        logs: List[str] = []
        for idx, row in enumerate(rows, start=2):
            result = await process_row(page, idx, row, logs, entity_type)
            results.append(result)
        await browser.close()

    failed_rows = [r for r in results if not r["success"]]
    success_count = len(results) - len(failed_rows)
    return {
        "message": f"Hoàn thành: {success_count} thành công, {len(failed_rows)} lỗi.",
        "results": results,
        "failed_rows": [r["row"] for r in failed_rows],
        "logs": logs,
    }


def import_excel_data(file_path: Path, entity_type: str):
    return asyncio.run(run_import(file_path, entity_type))
